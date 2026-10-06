import io
import streamlit as st
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from copy import copy

# Configuración de página
st.set_page_config(
    page_title="Prorrateo Terapias e Insumos",
    page_icon="📋",
    layout="wide"
)

# Etiqueta para prevenir que el traductor automático rompa la vista
st.markdown('<div translate="no">', unsafe_allow_html=True)

st.title("📋 Conciliación Mensual de Insumos y Terapias")
st.markdown("Carga los dos archivos de Excel para realizar el cruce automático, prorrateo y asignación de costos.")

# Panel lateral para carga de archivos
with st.sidebar:
    st.header("📂 Carga de Archivos")
    file_terapias = st.file_uploader("1. Archivo de Terapias (.xlsx)", type=["xlsx"])
    file_insumos = st.file_uploader("2. Archivo de Insumos (.xlsx)", type=["xlsx"])

if file_terapias and file_insumos:
    try:
        # Cargar datos
        df_insumos = pd.read_excel(file_insumos, sheet_name='TERAPIAS ')
        df_td = pd.read_excel(file_terapias, sheet_name='TERAPIA DETALLE')

        # Detección de cédulas en insumos
        nit_insumos = set(df_insumos['NIT'].dropna().unique())
        df_td['DOC_CLEAN'] = pd.to_numeric(df_td['DOCUMENTO CONTRATISTA'], errors='coerce')
        nit_terapias = set(df_td['DOC_CLEAN'].dropna().unique())

        novedades = nit_insumos - nit_terapias
        reasignaciones = {}

        if novedades:
            st.warning("⚠️ Se encontraron contratistas con insumos que no tienen terapias registradas este mes.")
            with st.expander("Revisar / Reasignar contratistas", expanded=True):
                for nit in novedades:
                    fila_n = df_insumos[df_insumos['NIT'] == nit].iloc[0]
                    costo_n = df_insumos[df_insumos['NIT'] == nit]['COSTO TOTAL '].sum()
                    
                    st.write(f"**Cédula:** `{nit}` — **Nombre:** `{fila_n['NOMBRE']}` | **Insumos:** `${costo_n:,.2f}`")
                    
                    c_nit, c_nom = st.columns(2)
                    with c_nit:
                        val_nit_def = "11428185" if nit == 38242451 else ""
                        nuevo_nit = st.text_input(f"Cédula destino para {fila_n['NOMBRE']}:", value=val_nit_def, key=f"n_{nit}")
                    with c_nom:
                        val_nom_def = "DONALDO RODRIGUEZ HERNANDEZ" if nit == 38242451 else ""
                        nuevo_nom = st.text_input(f"Nombre contratista destino:", value=val_nom_def, key=f"m_{nit}")
                    
                    if nuevo_nit.strip():
                        try:
                            reasignaciones[int(nit)] = (int(nuevo_nit.strip()), nuevo_nom.strip().upper())
                        except ValueError:
                            st.error(f"La cédula {nuevo_nit} debe contener solo números.")

        # Botón para procesar
        if st.button("🚀 Procesar Archivo", type="primary"):
            with st.spinner("Realizando cruce y prorrateo..."):
                # Aplicar reasignaciones
                for nit_orig, (nit_dest, nom_dest) in reasignaciones.items():
                    m = df_insumos['NIT'] == nit_orig
                    df_insumos.loc[m, 'NIT'] = nit_dest
                    df_insumos.loc[m, 'NOMBRE'] = nom_dest

                # Consolidar insumos por contratista
                resumen = df_insumos.groupby('NIT', as_index=False).agg({
                    'COSTO TOTAL ': 'sum',
                    'NOMBRE': 'first'
                }).rename(columns={'NIT': 'CEDULA', 'COSTO TOTAL ': 'VALOR TOTAL', 'NOMBRE': 'CONTRATISTA'})

                val_dict = dict(zip(resumen['CEDULA'], resumen['VALOR TOTAL']))

                # Cantidad total de terapias por contratista
                sum_terapias = df_td.groupby('DOC_CLEAN')['CANTIDAD AUTORIZADA'].sum().to_dict()

                # Valor unitario por terapia
                unit_val_dict = {}
                for doc, val in val_dict.items():
                    total_t = sum_terapias.get(doc, 0)
                    unit_val_dict[doc] = (val / total_t) if total_t > 0 else 0.0

                # Modificar archivo original con openpyxl
                file_terapias.seek(0)
                wb = openpyxl.load_workbook(file_terapias)

                # 1. Crear pestaña VALOR POR CONTRATISTA TERAPIAS
                if 'VALOR POR CONTRATISTA TERAPIAS' in wb.sheetnames:
                    del wb['VALOR POR CONTRATISTA TERAPIAS']
                ws_val = wb.create_sheet('VALOR POR CONTRATISTA TERAPIAS')

                headers_val = ['CEDULA', 'VALOR TOTAL', 'CONTRATISTA']
                h_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
                h_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
                border_thin = Border(
                    left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
                    top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9')
                )

                for col_idx, h in enumerate(headers_val, 1):
                    cell = ws_val.cell(1, col_idx, h)
                    cell.font = h_font
                    cell.fill = h_fill
                    cell.alignment = Alignment(horizontal='center', vertical='center')

                for i, row in resumen.iterrows():
                    r = i + 2
                    c1 = ws_val.cell(r, 1, row['CEDULA'])
                    c2 = ws_val.cell(r, 2, row['VALOR TOTAL'])
                    c3 = ws_val.cell(r, 3, row['CONTRATISTA'])
                    c1.alignment = Alignment(horizontal='center')
                    c2.number_format = '"$"#,##0.00'
                    c2.alignment = Alignment(horizontal='right')
                    c3.alignment = Alignment(horizontal='left')
                    for c in (c1, c2, c3):
                        c.border = border_thin

                ws_val.column_dimensions['A'].width = 18
                ws_val.column_dimensions['B'].width = 20
                ws_val.column_dimensions['C'].width = 42

                # 2. Modificar pestaña TERAPIA DETALLE
                ws_td = wb['TERAPIA DETALLE']
                ws_td.insert_cols(13) # Inserta TOTAL
                ws_td.insert_cols(22) # Inserta Columna8

                # Estilos encabezados
                h_total = ws_td.cell(1, 13, 'TOTAL')
                ref_12 = ws_td.cell(1, 12)
                if ref_12.has_style:
                    h_total.font = copy(ref_12.font)
                    h_total.fill = copy(ref_12.fill)
                    h_total.border = copy(ref_12.border)
                    h_total.alignment = copy(ref_12.alignment)

                h_col8 = ws_td.cell(1, 22, 'Columna8')
                ref_23 = ws_td.cell(1, 23)
                if ref_23.has_style:
                    h_col8.font = copy(ref_23.font)
                    h_col8.fill = copy(ref_23.fill)
                    h_col8.border = copy(ref_23.border)
                    h_col8.alignment = copy(ref_23.alignment)

                seen_docs = set()
                total_prorrateado = 0

                for r_idx in range(2, ws_td.max_row + 1):
                    cant = ws_td.cell(r_idx, 12).value
                    doc = ws_td.cell(r_idx, 23).value

                    try:
                        doc_clean = int(float(str(doc).strip()))
                    except (ValueError, TypeError):
                        doc_clean = None

                    try:
                        cant_clean = float(cant)
                    except (ValueError, TypeError):
                        cant_clean = 0.0

                    # Valor TOTAL fila
                    u_val = unit_val_dict.get(doc_clean, 0.0)
                    fila_total = round(u_val * cant_clean, 2)
                    cell_tot = ws_td.cell(r_idx, 13, fila_total)
                    cell_tot.number_format = '"$"#,##0.00'
                    cell_tot.alignment = Alignment(horizontal='right')
                    total_prorrateado += fila_total

                    # Valor Columna8
                    cell_c8 = ws_td.cell(r_idx, 22)
                    if doc_clean in val_dict:
                        if doc_clean not in seen_docs:
                            seen_docs.add(doc_clean)
                            c8_val = round(val_dict[doc_clean], 2)
                            cell_c8.value = c8_val
                            cell_c8.number_format = '"$"#,##0.00'
                            cell_c8.alignment = Alignment(horizontal='right')
                        else:
                            cell_c8.value = None
                    else:
                        cell_c8.value = None

                ws_td.column_dimensions['M'].width = 16
                ws_td.column_dimensions['V'].width = 18

                # Guardar en buffer en memoria
                out_buffer = io.BytesIO()
                wb.save(out_buffer)
                out_buffer.seek(0)

                st.success("✅ ¡Archivo procesado exitosamente!")

                # Métricas
                m1, m2, m3 = st.columns(3)
                m1.metric("Total Insumos", f"${resumen['VALOR TOTAL'].sum():,.2f}")
                m2.metric("Total Prorrateado", f"${total_prorrateado:,.2f}")
                m3.metric("Contratistas Conciliados", f"{len(seen_docs)} / {len(resumen)}")

                # Botón de Descarga
                nombre_salida = file_terapias.name.replace(".xlsx", "_CONCILIADO.xlsx")
                st.download_button(
                    label="📥 Descargar Archivo Excel Conciliado",
                    data=out_buffer,
                    file_name=nombre_salida,
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )

    except Exception as e:
        st.error(f"Error procesando los datos: {str(e)}")
else:
    st.info("👈 Sube ambos archivos en el menú de la izquierda para comenzar.")

st.markdown('</div>', unsafe_allow_html=True)